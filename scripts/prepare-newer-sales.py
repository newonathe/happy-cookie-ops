#!/usr/bin/env python3 -I
"""B4b: her Sales rows NEWER than what the trial's ledger already holds, as a ledger plan for scripts/prepare-newer-sales.mjs.

  python3 -I scripts/prepare-newer-sales.py --workbook <xlsx> --after <YYYY-MM-DD> --out <dir>

It reads one file, writes only to --out (ledger_plan.json, report.json) and never connects to a database. Keep --out outside Git; the
plan holds her sale amounts. Only counts, sums, dates and hashes are meant to be pasted anywhere. `--after` is the trial's last imported
sale day, which Ethan reads with scripts/b4b-trial-reads.sql.

The rules are the ones Ethan approved for the original 804-line import (docs/state/F2-real-data.md, "Historical financial import", 5 October):
  1. a sale row with no date inherits the preceding explicit date (the row above it, in sheet order);
  2. repeated identical lines are separate genuine sales (every row is its own source ref, so none is merged);
  3. Asia/Manila is +08:00 (no daylight saving); a date is midnight that day, as in the original rows (tests/unit/historical-ledger-plan.test.ts);
  4. a row without a valid amount is held, not written (28 such rows were held in the original import); "Sales (GForm)" is never counted.
Anything those rules do not settle is LISTED under `unresolved`, never guessed: no date above the first row, a date cell that is not a date (and
the blank rows below it), a future date, an amount that is missing, text, zero, negative or has a fraction of a centavo. The memo is
"Workbook Sales row N": no customer, note or payment text is carried. A row dated on or before --after but sitting below newer rows is
listed under `to_confirm` (it may be a late entry that was never imported); it is never planned.
Ids are uuid5(source ref + source hash), so a re-run gives the same plan and the same hash.
"""
import argparse
import datetime
import hashlib
import json
import sys
import uuid
from decimal import Decimal
from pathlib import Path

import openpyxl

NAMESPACE = uuid.uuid5(uuid.NAMESPACE_DNS, 'historical-sales.happy-cookie-ops.invalid')
MANILA = datetime.timezone(datetime.timedelta(hours=8))
RULES = [
    'a blank date inherits the preceding explicit date',
    'repeated identical lines are separate sales',
    'each date is midnight Asia/Manila (+08:00)',
    'a row without a valid amount is held, listed, not written',
]


def parse_day(s, what):
    try:
        return datetime.date.fromisoformat(str(s))
    except ValueError:
        sys.exit('%s must be a calendar date YYYY-MM-DD, got %r' % (what, s))


def day_of(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return None


def cents_of(v):
    """(centavos, why-not)"""
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return None, 'no usable amount'
    pesos = Decimal(str(v))
    if pesos <= 0:
        return None, 'no usable amount (zero or negative)'
    c = pesos * 100
    if c != c.to_integral_value():
        return None, 'amount has a fraction of a centavo'
    return int(c), None


def build(path, after, source_sha256, today=None):
    path = Path(path)
    after = parse_day(after, '--after')
    today = today or datetime.datetime.now(MANILA).date()
    wb = openpyxl.load_workbook(path, data_only=True)
    if 'dummy' in path.name.lower() or any('dummy' in n.lower() for n in wb.sheetnames):
        sys.exit('this looks like the DUMMY workbook; it is excluded entirely (Ethan, 5 October)')
    if 'Sales' not in wb.sheetnames:
        sys.exit('no Sales sheet')
    plan, unresolved, to_confirm = [], [], []
    eff, blocked, seen_newer, on_after, total = None, False, False, 0, 0
    for n, r in enumerate(wb['Sales'].iter_rows(min_row=2, max_col=9, values_only=True), start=2):
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in r):
            continue
        total += 1
        source = 'Sales!A%d:I%d' % (n, n)
        raw = r[0]
        date_why = None
        if raw is not None and not (isinstance(raw, str) and not raw.strip()):
            d = day_of(raw)
            if d is None:
                eff, blocked, date_why = None, True, 'the date cell is not a date'
            else:
                eff, blocked = d, False
        elif blocked:
            date_why = 'follows an unreadable date; the date above cannot be inherited'
        elif eff is None:
            date_why = 'no date, and none above to inherit'
        if date_why is None and eff > today:
            date_why = 'date is in the future'
        if date_why:
            unresolved.append(dict(row=n, source=source, why=date_why))
            continue
        cents, amount_why = cents_of(r[6])
        if eff == after:
            on_after += 1
        if eff <= after:
            if seen_newer and cents is not None:
                to_confirm.append(dict(row=n, source=source, date=eff.isoformat(),
                                       why='dated on or before --after but sits below newer rows: confirm it is already in the ledger'))
            continue
        seen_newer = True
        if amount_why:
            unresolved.append(dict(row=n, source=source, why=amount_why))
            continue
        plan.append(dict(id=str(uuid.uuid5(NAMESPACE, source + source_sha256)), kind='sale', category='Historical sales',
                         memo='Workbook Sales row %d' % n, amount_cents=cents, occurred_at='%sT00:00:00+08:00' % eff.isoformat(), source=source))
    canon = json.dumps(plan, separators=(',', ':'), ensure_ascii=False)
    dates = sorted(p['occurred_at'][:10] for p in plan)
    report = dict(source_sha256=source_sha256, after=after.isoformat(),
                  counts=dict(sales_rows=total, planned=len(plan), unresolved=len(unresolved), on_after_date=on_after, to_confirm=len(to_confirm)),
                  sales_cents=sum(p['amount_cents'] for p in plan), first_date=dates[0] if dates else None, last_date=dates[-1] if dates else None,
                  plan_sha256=hashlib.sha256(canon.encode()).hexdigest(), rules=RULES, unresolved=unresolved, to_confirm=to_confirm)
    return plan, report


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument('--workbook', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--after', required=True, help="the trial's last imported sale day (YYYY-MM-DD); only later days are planned")
    ap.add_argument('--today', help='override today (YYYY-MM-DD), for tests')
    a = ap.parse_args(argv)
    after = parse_day(a.after, '--after')
    today = parse_day(a.today, '--today') if a.today else None
    source = hashlib.sha256(Path(a.workbook).read_bytes()).hexdigest()
    plan, report = build(a.workbook, after.isoformat(), source, today)
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    (out / 'ledger_plan.json').write_text(json.dumps(dict(source_sha256=source, dummy_workbook_excluded_entirely=True, ledger_plan=plan),
                                                     indent=1, ensure_ascii=False))
    (out / 'report.json').write_text(json.dumps(report, indent=1, ensure_ascii=False))
    print(json.dumps({k: report[k] for k in ('counts', 'sales_cents', 'first_date', 'last_date', 'plan_sha256', 'source_sha256')}, indent=1))


if __name__ == '__main__':
    main()
